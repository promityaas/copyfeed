import openpyxl
p='artifacts/export-validation.xlsx'
w=openpyxl.load_workbook(p)
assert w.sheetnames==['Posts and metrics','Outlier rankings','Hook analysis','Content pillars','Audience questions','Opportunity Scores','Research notes']
s=w['Posts and metrics']
headers={c.value:c.column for c in s[1]}
assert s.cell(2,headers['title']).value=='=1+1'
assert s.cell(2,headers['title']).data_type=='s'
assert s.cell(2,headers['caption']).value=='हिन्दी'
assert s.cell(2,headers['views']).value is None
print('XLSX independently parsed: seven sheets, Unicode, literal formula text, missing views blank.')
